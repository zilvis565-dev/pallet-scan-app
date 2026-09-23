from flask import Flask, render_template, request, redirect, url_for, session, send_file
import sqlite3
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.secret_key = "supersecretkey"

DB_NAME = "data.db"


def get_conn():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_conn()
    c = conn.cursor()

    c.execute("""
        CREATE TABLE IF NOT EXISTS orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            order_number TEXT,
            created_at TEXT NOT NULL
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS pallets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            order_id INTEGER NOT NULL,
            pallet_number INTEGER NOT NULL,
            pallet_code TEXT NOT NULL,
            length INTEGER,
            width INTEGER,
            height INTEGER,
            weight REAL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(order_id) REFERENCES orders(id)
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS pallet_lines (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pallet_id INTEGER NOT NULL,
            part_code TEXT NOT NULL,
            description TEXT,
            pcs INTEGER NOT NULL,
            boxes_qty INTEGER NOT NULL,
            pcs_per_box INTEGER NOT NULL,
            FOREIGN KEY(pallet_id) REFERENCES pallets(id)
        )
    """)

    conn.commit()
    conn.close()


def generate_pallet_code():
    conn = get_conn()
    today_compact = datetime.now().strftime("%Y%m%d")
    today_db = datetime.now().strftime("%Y-%m-%d")

    result = conn.execute("""
        SELECT COUNT(*) as cnt
        FROM pallets
        WHERE substr(created_at, 1, 10) = ?
    """, (today_db,)).fetchone()

    conn.close()

    next_number = result["cnt"] + 1
    return f"{today_compact}-{next_number:03d}"


def get_next_pallet_number(order_id):
    conn = get_conn()
    result = conn.execute("""
        SELECT MAX(pallet_number) as max_pallet
        FROM pallets
        WHERE order_id = ?
    """, (order_id,)).fetchone()
    conn.close()

    if result["max_pallet"] is None:
        return 1
    return result["max_pallet"] + 1


@app.route("/")
def dashboard():
    conn = get_conn()
    orders = conn.execute("""
        SELECT * FROM orders
        ORDER BY id DESC
    """).fetchall()
    conn.close()
    return render_template("dashboard.html", orders=orders)


@app.route("/new-order", methods=["GET", "POST"])
def new_order():
    if request.method == "POST":
        order_number = request.form.get("order_number", "").strip()

        conn = get_conn()
        c = conn.cursor()
        c.execute("""
            INSERT INTO orders (order_number, created_at)
            VALUES (?, ?)
        """, (
            order_number if order_number else None,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ))
        order_id = c.lastrowid
        conn.commit()
        conn.close()

        session["order_id"] = order_id
        session["order_number"] = order_number if order_number else ""
        session["current_pallet"] = get_next_pallet_number(order_id)
        session["current_lines"] = {}
        session["mixed_box_lines"] = []

        return redirect(url_for("scan"))

    return render_template("new_order.html")


@app.route("/delete-order/<int:order_id>", methods=["POST"])
def delete_order(order_id):
    conn = get_conn()

    pallet_ids = conn.execute("""
        SELECT id FROM pallets WHERE order_id = ?
    """, (order_id,)).fetchall()

    for pallet in pallet_ids:
        conn.execute("""
            DELETE FROM pallet_lines WHERE pallet_id = ?
        """, (pallet["id"],))

    conn.execute("""
        DELETE FROM pallets WHERE order_id = ?
    """, (order_id,))

    conn.execute("""
        DELETE FROM orders WHERE id = ?
    """, (order_id,))

    conn.commit()
    conn.close()

    return redirect(url_for("dashboard"))


@app.route("/delete-pallet/<int:pallet_id>", methods=["POST"])
def delete_pallet(pallet_id):
    conn = get_conn()

    pallet = conn.execute("""
        SELECT * FROM pallets WHERE id = ?
    """, (pallet_id,)).fetchone()

    if not pallet:
        conn.close()
        return redirect(url_for("dashboard"))

    order_id = pallet["order_id"]

    conn.execute("""
        DELETE FROM pallet_lines WHERE pallet_id = ?
    """, (pallet_id,))

    conn.execute("""
        DELETE FROM pallets WHERE id = ?
    """, (pallet_id,))

    conn.commit()
    conn.close()

    return redirect(url_for("order_detail", order_id=order_id))


@app.route("/order/<int:order_id>")
def order_detail(order_id):
    conn = get_conn()

    order = conn.execute("""
        SELECT * FROM orders WHERE id = ?
    """, (order_id,)).fetchone()

    pallets = conn.execute("""
        SELECT * FROM pallets
        WHERE order_id = ?
        ORDER BY pallet_number
    """, (order_id,)).fetchall()

    pallet_data = []
    total_pcs = 0
    total_boxes = 0
    total_weight = 0

    for pallet in pallets:
        lines = conn.execute("""
            SELECT * FROM pallet_lines
            WHERE pallet_id = ?
            ORDER BY id
        """, (pallet["id"],)).fetchall()

        pallet_data.append({"pallet": pallet, "lines": lines})

        total_weight += pallet["weight"] if pallet["weight"] else 0
        for line in lines:
            total_pcs += line["pcs"]
            total_boxes += line["boxes_qty"]

    conn.close()
    return render_template(
        "order_detail.html",
        order=order,
        pallet_data=pallet_data,
        total_pcs=total_pcs,
        total_boxes=total_boxes,
        total_weight=total_weight
    )


@app.route("/start-pallet/<int:order_id>")
def start_pallet(order_id):
    conn = get_conn()
    order = conn.execute("""
        SELECT * FROM orders WHERE id = ?
    """, (order_id,)).fetchone()
    conn.close()

    session["order_id"] = order["id"]
    session["order_number"] = order["order_number"] if order["order_number"] else ""
    session["current_pallet"] = get_next_pallet_number(order_id)
    session["current_lines"] = {}
    session["mixed_box_lines"] = []

    return redirect(url_for("scan"))


@app.route("/scan", methods=["GET", "POST"])
def scan():
    if "order_id" not in session:
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        part_code = request.form["part_code"].strip().lower()
        pcs_per_box = int(request.form["pcs_per_box"])

        if part_code:
            lines = session.get("current_lines", {})

            if part_code in lines:
                lines[part_code]["pcs"] += pcs_per_box
                lines[part_code]["boxes_qty"] += 1
            else:
                lines[part_code] = {
                    "description": part_code,
                    "pcs": pcs_per_box,
                    "boxes_qty": 1,
                    "pcs_per_box": pcs_per_box
                }

            session["current_lines"] = lines

        return redirect(url_for("scan"))

    lines = session.get("current_lines", {})
    total_pcs = sum(item["pcs"] for item in lines.values())
    total_boxes = sum(item["boxes_qty"] for item in lines.values())

    return render_template(
        "scan.html",
        order_number=session["order_number"],
        current_pallet=session["current_pallet"],
        lines=lines,
        total_pcs=total_pcs,
        total_boxes=total_boxes
    )


@app.route("/mixed-box", methods=["GET", "POST"])
def mixed_box():
    if "order_id" not in session:
        return redirect(url_for("dashboard"))

    if "mixed_box_lines" not in session:
        session["mixed_box_lines"] = []

    if request.method == "POST":
        action = request.form.get("action")

        if action == "add":
            part_code = request.form["part_code"].strip().lower()
            pcs = int(request.form["pcs"])

            if part_code:
                mixed_lines = session.get("mixed_box_lines", [])
                mixed_lines.append({
                    "part_code": part_code,
                    "description": part_code,
                    "pcs": pcs
                })
                session["mixed_box_lines"] = mixed_lines

            return redirect(url_for("mixed_box"))

        elif action == "save":
            mixed_lines = session.get("mixed_box_lines", [])
            current_lines = session.get("current_lines", {})

            for idx, line in enumerate(mixed_lines):
                boxes_to_add = 1 if idx == 0 else 0

                if line["part_code"] in current_lines:
                    current_lines[line["part_code"]]["pcs"] += line["pcs"]
                    current_lines[line["part_code"]]["boxes_qty"] += boxes_to_add
                else:
                    current_lines[line["part_code"]] = {
                        "description": line["description"],
                        "pcs": line["pcs"],
                        "boxes_qty": boxes_to_add,
                        "pcs_per_box": line["pcs"]
                    }

            session["current_lines"] = current_lines
            session["mixed_box_lines"] = []
            return redirect(url_for("scan"))

        elif action == "cancel":
            session["mixed_box_lines"] = []
            return redirect(url_for("scan"))

    return render_template(
        "mixed_box.html",
        order_number=session["order_number"],
        current_pallet=session["current_pallet"],
        mixed_box_lines=session.get("mixed_box_lines", [])
    )


@app.route("/finish-pallet", methods=["GET", "POST"])
def finish_pallet():
    if "order_id" not in session:
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        length = int(request.form["length"])
        width = int(request.form["width"])
        height = int(request.form["height"])
        weight = float(request.form["weight"])

        pallet_code = generate_pallet_code()

        conn = get_conn()
        c = conn.cursor()

        c.execute("""
            INSERT INTO pallets (
                order_id, pallet_number, pallet_code,
                length, width, height, weight, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            session["order_id"],
            session["current_pallet"],
            pallet_code,
            length,
            width,
            height,
            weight,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ))

        pallet_id = c.lastrowid

        for part_code, data in session["current_lines"].items():
            c.execute("""
                INSERT INTO pallet_lines (
                    pallet_id, part_code, description,
                    pcs, boxes_qty, pcs_per_box
                ) VALUES (?, ?, ?, ?, ?, ?)
            """, (
                pallet_id,
                part_code,
                data["description"],
                data["pcs"],
                data["boxes_qty"],
                data["pcs_per_box"]
            ))

        conn.commit()
        conn.close()

        order_id = session["order_id"]
        session.pop("current_lines", None)
        session.pop("current_pallet", None)
        session.pop("order_number", None)
        session.pop("order_id", None)
        session.pop("mixed_box_lines", None)

        return redirect(url_for("order_detail", order_id=order_id))

    return render_template("edit_pallet.html", mode="new")


@app.route("/edit-pallet/<int:pallet_id>", methods=["GET", "POST"])
def edit_pallet(pallet_id):
    conn = get_conn()

    pallet = conn.execute("""
        SELECT * FROM pallets WHERE id = ?
    """, (pallet_id,)).fetchone()

    if request.method == "POST":
        length = int(request.form["length"])
        width = int(request.form["width"])
        height = int(request.form["height"])
        weight = float(request.form["weight"])

        conn.execute("""
            UPDATE pallets
            SET length = ?, width = ?, height = ?, weight = ?
            WHERE id = ?
        """, (length, width, height, weight, pallet_id))
        conn.commit()
        conn.close()
        return redirect(url_for("order_detail", order_id=pallet["order_id"]))

    conn.close()
    return render_template("edit_pallet.html", mode="edit", pallet=pallet)


@app.route("/edit-line/<int:line_id>", methods=["POST"])
def edit_line(line_id):
    pcs = int(request.form["pcs"])
    boxes_qty = int(request.form["boxes_qty"])
    description = request.form["description"]

    conn = get_conn()

    line = conn.execute("""
        SELECT pallet_lines.*, pallets.order_id
        FROM pallet_lines
        JOIN pallets ON pallet_lines.pallet_id = pallets.id
        WHERE pallet_lines.id = ?
    """, (line_id,)).fetchone()

    conn.execute("""
        UPDATE pallet_lines
        SET pcs = ?, boxes_qty = ?, description = ?
        WHERE id = ?
    """, (pcs, boxes_qty, description, line_id))

    conn.commit()
    conn.close()

    return redirect(url_for("order_detail", order_id=line["order_id"]))


@app.route("/delete-line/<int:line_id>", methods=["POST"])
def delete_line(line_id):
    conn = get_conn()

    line = conn.execute("""
        SELECT pallet_lines.*, pallets.order_id
        FROM pallet_lines
        JOIN pallets ON pallet_lines.pallet_id = pallets.id
        WHERE pallet_lines.id = ?
    """, (line_id,)).fetchone()

    conn.execute("DELETE FROM pallet_lines WHERE id = ?", (line_id,))
    conn.commit()
    conn.close()

    return redirect(url_for("order_detail", order_id=line["order_id"]))


@app.route("/export-order/<int:order_id>")
def export_order(order_id):
    conn = get_conn()

    order = conn.execute("""
        SELECT * FROM orders WHERE id = ?
    """, (order_id,)).fetchone()

    pallets = conn.execute("""
        SELECT * FROM pallets
        WHERE order_id = ?
        ORDER BY pallet_number
    """, (order_id,)).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "PackingList"

    thin = Side(style="thin")
    medium = Side(style="medium")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    total_border = Border(left=medium, right=medium, top=medium, bottom=medium)

    header_fill = PatternFill(fill_type="solid", start_color="3B82F6", end_color="3B82F6")
    header_font = Font(bold=True, color="FFFFFF")
    bold_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A1"] = "Shipment date"
    ws["A1"].font = bold_font
    ws["B1"] = datetime.now().strftime("%Y-%m-%d")

    headers = [
        "Pallet Nr.",
        "Pallet code",
        "Order number (reference)",
        "Code",
        "PCS",
        "Boxes qty",
        "Pallet measurements",
        "Pallet weight"
    ]

    header_row = 3
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    current_row = 4
    grand_total_pcs = 0
    grand_total_boxes = 0
    grand_total_weight = 0

    for pallet in pallets:
        lines = conn.execute("""
            SELECT * FROM pallet_lines
            WHERE pallet_id = ?
            ORDER BY id
        """, (pallet["id"],)).fetchall()

        if not lines:
            continue

        start_row = current_row
        end_row = current_row + len(lines) - 1

        total_boxes_qty = sum(line["boxes_qty"] for line in lines)
        measurements = f'{pallet["length"]}*{pallet["width"]}*{pallet["height"]}'

        grand_total_boxes += total_boxes_qty
        grand_total_weight += pallet["weight"]

        for i, line in enumerate(lines):
            row = current_row + i

            ws.cell(row=row, column=4, value=line["description"])
            ws.cell(row=row, column=5, value=line["pcs"])

            grand_total_pcs += line["pcs"]

            ws.cell(row=row, column=4).alignment = center
            ws.cell(row=row, column=5).alignment = center
            ws.cell(row=row, column=4).border = border
            ws.cell(row=row, column=5).border = border

        ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
        ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
        ws.merge_cells(start_row=start_row, start_column=3, end_row=end_row, end_column=3)
        ws.merge_cells(start_row=start_row, start_column=6, end_row=end_row, end_column=6)
        ws.merge_cells(start_row=start_row, start_column=7, end_row=end_row, end_column=7)
        ws.merge_cells(start_row=start_row, start_column=8, end_row=end_row, end_column=8)

        ws.cell(row=start_row, column=1, value=pallet["pallet_number"])
        ws.cell(row=start_row, column=2, value=pallet["pallet_code"])
        ws.cell(row=start_row, column=3, value=order["order_number"] if order["order_number"] else "")
        ws.cell(row=start_row, column=6, value=total_boxes_qty)
        ws.cell(row=start_row, column=7, value=measurements)
        ws.cell(row=start_row, column=8, value=pallet["weight"])

        for col in [1, 2, 3, 6, 7, 8]:
            cell = ws.cell(row=start_row, column=col)
            cell.alignment = center
            cell.border = border

        for row in range(start_row, end_row + 1):
            for col in range(1, 9):
                ws.cell(row=row, column=col).border = border

        current_row = end_row + 1

    total_row = current_row + 2

    ws.cell(row=total_row, column=4, value="TOTALS")
    ws.cell(row=total_row, column=5, value=grand_total_pcs)
    ws.cell(row=total_row, column=6, value=grand_total_boxes)
    ws.cell(row=total_row, column=8, value=grand_total_weight)

    for col in [4, 5, 6, 8]:
        ws.cell(row=total_row, column=col).font = bold_font
        ws.cell(row=total_row, column=col).border = total_border
        ws.cell(row=total_row, column=col).alignment = center

    summary_ws = wb.create_sheet(title="Summary")

    summary_headers = ["Code", "Description", "Total PCS", "Total Boxes qty"]
    for col_num, header in enumerate(summary_headers, 1):
        cell = summary_ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    summary_rows = conn.execute("""
        SELECT
            part_code,
            description,
            SUM(pcs) AS total_pcs,
            SUM(boxes_qty) AS total_boxes_qty
        FROM pallet_lines
        JOIN pallets ON pallet_lines.pallet_id = pallets.id
        WHERE pallets.order_id = ?
        GROUP BY part_code, description
        ORDER BY description
    """, (order_id,)).fetchall()

    row_num = 2
    for row in summary_rows:
        summary_ws.cell(row=row_num, column=1, value=row["part_code"])
        summary_ws.cell(row=row_num, column=2, value=row["description"])
        summary_ws.cell(row=row_num, column=3, value=row["total_pcs"])
        summary_ws.cell(row=row_num, column=4, value=row["total_boxes_qty"])

        for col in range(1, 5):
            summary_ws.cell(row=row_num, column=col).border = border
            summary_ws.cell(row=row_num, column=col).alignment = center

        row_num += 1

    summary_widths = {
        1: 18,
        2: 28,
        3: 14,
        4: 16
    }

    for col_idx, width in summary_widths.items():
        summary_ws.column_dimensions[get_column_letter(col_idx)].width = width

    summary_ws.freeze_panes = "A2"

    conn.close()

    widths = {
        1: 12,
        2: 22,
        3: 22,
        4: 30,
        5: 10,
        6: 12,
        7: 22,
        8: 14
    }

    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = 24

    filename = f"packing_list_order_{order['id']}.xlsx"
    wb.save(filename)

    return send_file(filename, as_attachment=True)


# IMPORTANT for Render / Waitress
init_db()

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
